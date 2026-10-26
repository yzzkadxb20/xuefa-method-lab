# -*- coding: utf-8 -*-
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE = Path(r"D:\deepseek-harness\default-workspace\deliverables\assets\base5.pptx")
prs = Presentation(str(BASE))
s = prs.slides[0]

MINT = RGBColor(0xB9, 0xF2, 0xD0)

for shape in list(s.shapes):
    if not shape.has_text_frame:
        continue
    txt = shape.text_frame.text.strip()
    if "MEMORY" in txt:
        shape.left = Inches(9.45)
        shape.top = Inches(5.68)
        shape.width = Inches(3.1)
        shape.height = Inches(0.24)
        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "MEMORY · 记忆贯穿始终"
        r.font.name = "Microsoft YaHei"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = MINT
    elif txt == "第二届全国人工智能应用创新大赛 · 智能体应用" and shape.top < Inches(7.0):
        shape._element.getparent().remove(shape._element)
    elif "2026.10" in txt and "METHOD LAB" in txt:
        shape.top = Inches(6.64)

# slide 5: move the bottom "why" strip down so it no longer touches the third card
s5 = prs.slides[4]
for shape in list(s5.shapes):
    if shape.has_text_frame and shape.text_frame.text.strip().startswith("为什么值得做"):
        shape.top = Inches(6.36)
    elif abs(shape.top.inches - 6.05) < 0.03 and abs(shape.width.inches - 11.9) < 0.15:
        shape.top = Inches(6.24)

prs.save(str(BASE))
print("patched", BASE)
