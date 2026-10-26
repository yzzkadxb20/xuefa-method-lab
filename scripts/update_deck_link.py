# -*- coding: utf-8 -*-
"""Replace the demo-link placeholder on slide 19 with a real public URL.

Usage:
    python scripts/update_deck_link.py https://your-site.netlify.app

The URL is also added as a clickable hyperlink. Keep the PPT filename unchanged
unless you also update the submission naming.
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

DECK = Path(r"D:\deepseek-harness\default-workspace\deliverables\学法实验室-参赛PPT.pptx")
CREAM = RGBColor(0xFF, 0xFD, 0xF8)


def main():
    if len(sys.argv) < 2:
        print("usage: python scripts/update_deck_link.py https://your-site.netlify.app [deck.pptx]")
        return 1
    url = sys.argv[1].strip()
    deck = Path(sys.argv[2]) if len(sys.argv) > 2 else DECK
    if not (url.startswith("http://") or url.startswith("https://")):
        print("URL must start with http:// or https://")
        return 1
    prs = Presentation(str(deck))
    slide = prs.slides[18]
    changed = 0
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text
        if "提交前替换" in text:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = f"{url}  ·  demo 账号：体验用户（无需密码）"
            r.font.name = "Microsoft YaHei"
            r.font.size = Pt(8)
            r.font.color.rgb = CREAM
            r.hyperlink.address = url
            changed += 1
    if not changed:
        print("placeholder not found; deck may already be updated")
        return 2
    prs.save(str(deck))
    print(f"updated slide 19 with: {url}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
