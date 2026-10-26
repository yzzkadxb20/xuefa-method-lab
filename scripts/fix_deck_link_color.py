# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pptx.oxml.ns import qn

DECK = r"D:\deepseek-harness\default-workspace\deliverables\学法实验室-参赛PPT.pptx"
URL = "https://roaring-pixie-05f195.netlify.app"
CREAM = RGBColor(0xE8, 0xF7, 0xEE)

prs = Presentation(DECK)
slide = prs.slides[18]
fixed = 0
for shape in slide.shapes:
    if not shape.has_text_frame:
        continue
    if "roaring-pixie" not in shape.text_frame.text:
        continue
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            # remove text-level hyperlink (it forces the dark blue theme color)
            rPr = r._r.get_or_add_rPr()
            for hl in list(rPr.findall(qn("a:hlinkClick"))):
                rPr.remove(hl)
            r.font.color.rgb = CREAM
            r.font.size = Pt(8)
            r.font.bold = False
    # make the whole box clickable without recoloring the text
    try:
        shape.click_action.hyperlink.address = URL
    except Exception as e:
        print("click action failed:", e)
    fixed += 1

prs.save(DECK)
print("fixed", fixed)
